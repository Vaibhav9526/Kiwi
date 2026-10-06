# KIWI x SIH26159 "SecureMailScope" deck — built ON TOP of the official
# SIH2026-IDEA-Presentation-Format.pptx template (structure kept intact).
# Run: python presentation/make_sih_pptx.py [out.pptx]
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

ROOT = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(ROOT)
TPL = os.path.join(ROOT, "SIH2026-IDEA-Presentation-Format.pptx")
OUT = os.path.join(ROOT, sys.argv[1] if len(sys.argv) > 1
                   else "KIWI-SIH26159-SecureMailScope.pptx")

SHOT_L = os.path.join(REPO, "docs", "screenshots", "mail-light.png")
SHOT_D = os.path.join(REPO, "docs", "screenshots", "mail-dark.png")
ARCH = os.path.join(ROOT, "kiwi-architecture.png")
LOGO = os.path.join(REPO, "images", "logo.png")

INK = RGBColor(0x10, 0x23, 0x1A)
ACCENT = RGBColor(0x1E, 0x8E, 0x3E)
ACCENT_D = RGBColor(0x14, 0x6A, 0x2E)
ACCENT_L = RGBColor(0x9F, 0xE6, 0xB8)
MUTED = RGBColor(0x5F, 0x6B, 0x66)
CARD = RGBColor(0xF1, 0xF7, 0xF3)
LINE = RGBColor(0xC9, 0xDC, 0xD0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB3, 0x2D, 0x2D)
AMBER = RGBColor(0xB8, 0x7A, 0x00)
FONT = "Arial"
FONT_T = "Times New Roman"

prs = Presentation(TPL)
W, H = prs.slide_width, prs.slide_height
S = list(prs.slides)

# Remove the template's prompt TextBox 8 on content slides — the pointer
# text is re-rendered verbatim by `pointers()` in a compact strip instead.
for s in S[1:6]:
    for sh in list(s.shapes):
        if sh.name == "TextBox 8" and sh.has_text_frame:
            sp = sh._element
            sp.getparent().remove(sp)


def box(s, x, y, w, h, fill=None, line=None, round_=True, shape=None):
    st = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE)
    sh = s.shapes.add_shape(st, x, y, w, h)
    if st == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = 0.08
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=4, line_spacing=1.05, font=FONT):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for txt, o in (para if isinstance(para, list) else [(para, {})]):
            r = p.add_run()
            r.text = txt
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def set_tf(shape, paras, font=FONT):
    """Rewrite a template text frame, preserving position/size."""
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for txt, o in para:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", 20))
            f.bold = o.get("bold", True)
            f.italic = o.get("italic", False)
            if o.get("color"):
                f.color.rgb = o["color"]


def find(s, name):
    for sh in s.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def pic(s, path, x, y, w=None, h=None, border=True):
    p = s.shapes.add_picture(path, x, y, width=w, height=h)
    if border:
        p.line.color.rgb = LINE
        p.line.width = Pt(1)
    p.shadow.inherit = False
    return p


def pointers(s, lines, y=Inches(1.22)):
    """Re-render the template's idea-detail pointers verbatim, small."""
    text(s, Inches(0.67), y, Inches(12.2), Inches(0.6),
         [[(l, {"size": 10.5, "italic": True, "color": MUTED, "bold": False})]
          for l in lines], space_after=1, line_spacing=1.0)


def bullets(s, x, y, w, h, items, size=13.5, gap=6):
    paras = []
    for head, body in items:
        segs = [("▪ ", {"color": ACCENT, "bold": True, "size": size})]
        if head:
            segs.append((head + " — ", {"bold": True, "color": ACCENT_D}))
        segs.append((body, {"color": INK}))
        paras.append(segs)
    return text(s, x, y, w, h, paras, size=size, space_after=gap)


def chip(s, x, y, w, h, label, fill=CARD, color=INK, size=11.5, bold=True, line=LINE):
    box(s, x, y, w, h, fill=fill, line=line)
    text(s, x, y, w, h, label, size=size, color=color, bold=bold,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)


# ============================================================ SLIDE 1 · TITLE
s = S[0]
set_tf(find(s, "Subtitle 3"),
       [[("KIWI  ·  SecureMailScope", {"font": FONT_T, "size": 30, "bold": True})]])
set_tf(find(s, "TextBox 9"), [
    [("Problem Statement ID – ", {}), ("SIH26159", {"color": ACCENT_D})],
    [("Problem Statement Title – SecureMailScope: AI-Assisted Cryptographic "
      "Security Posture Assessment for Secure Email Communications",
      {"size": 15})],
    [("Theme – ", {}), ("Blockchain & Cybersecurity", {"color": ACCENT_D})],
    [("PS Category – ", {}), ("Software", {"color": ACCENT_D})],
    [("Team ID – __________", {})],
    [("Team Name (Registered on portal) – __________", {"size": 18})],
])
pic(s, LOGO, Inches(0.42), Inches(6.28), w=Inches(0.62), border=False)
text(s, Inches(1.14), Inches(6.32), Inches(5.6), Inches(0.6), [[
    ("KIWI", {"bold": True, "size": 15, "color": ACCENT_D}),
    (" — email that proves its security.", {"size": 13, "color": MUTED}),
]])

# ====================================================== SLIDE 2 · IDEA TITLE
s = S[1]
pointers(s, ["Proposed Solution (Describe your Idea/Solution/Prototype) · "
             "Detailed explanation · How it addresses the problem · "
             "Innovation and uniqueness of the solution"])
text(s, Inches(0.67), Inches(1.5), Inches(12.2), Inches(0.55), [[
    ("KIWI — SecureMailScope  ", {"size": 26, "bold": True, "color": INK}),
    ("|  Passive PCAP forensics + endpoint TLS evidence for "
     "SMTP · IMAP · POP3", {"size": 15, "bold": True, "color": ACCENT_D}),
]])
bullets(s, Inches(0.67), Inches(2.2), Inches(7.4), Inches(2.2), [
    ("Proposed solution",
     "a passive forensic engine that ingests PCAP, rebuilds SMTP/IMAP/POP3 "
     "sessions and TLS handshakes, validates certificates, then scores the "
     "cryptographic posture."),
    ("How it addresses the problem",
     "deterministic rules flag weak TLS, broken STARTTLS and bad certs; an AI "
     "layer classifies risk, spots anomalous handshakes and recommends fixes — "
     "it explains, it never asserts."),
    ("Innovation & uniqueness",
     "the only entry that is also a working mail client — endpoint evidence "
     "sees what passive capture cannot (comparison below)."),
], size=13, gap=8)
pic(s, SHOT_L, Inches(8.35), Inches(1.98), w=Inches(4.3))

# ---- comparison block --------------------------------------------------
text(s, Inches(0.67), Inches(4.62), Inches(9), Inches(0.32), [[
    ("HOW KIWI DIFFERS FROM EXISTING TOOLS", {"size": 13, "bold": True,
                                             "color": ACCENT_D}),
]])
rows = [
    ("Capability", "Wireshark", "Zeek", "SSLyze / testssl.sh", "KIWI"),
    ("Passive PCAP analysis — no live probing", "✓", "✓", "✗ active scan only", "✓"),
    ("Email-aware: SMTP · IMAP · POP3 + STARTTLS", "~ decode only", "~ scripts needed", "✗", "✓ native"),
    ("Auto posture score + prioritised findings", "✗", "✗", "~ grade only", "✓ rule-based score"),
    ("AI risk classification + anomaly detection", "✗", "✗", "✗", "✓ explains, never asserts"),
    ("Endpoint vantage + self-verifying report", "✗", "✗", "✗", "✓ SHA-256 envelope"),
]
tx, ty, tw = Inches(0.67), Inches(4.98), Inches(12.28)
gtbl = s.shapes.add_table(len(rows), 5, tx, ty, tw, Inches(1.82)).table
gtbl.columns[0].width = Inches(5.0)
for c in range(1, 5):
    gtbl.columns[c].width = Inches(1.82)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = gtbl.cell(r, c)
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.01)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        run = p.add_run()
        run.text = val
        f = run.font
        f.name = FONT
        f.size = Pt(11 if r == 0 else 10)
        f.bold = (r == 0) or (c == 0)
        if r == 0:
            cell.fill.solid(); cell.fill.fore_color.rgb = ACCENT_D
            f.color.rgb = WHITE
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD if r % 2 else WHITE
            if c == 4:
                cell.fill.fore_color.rgb = RGBColor(0xDD, 0xF3, 0xE4)
                f.color.rgb = ACCENT_D
                f.bold = True
            elif val.startswith("✓"):
                f.color.rgb = ACCENT_D
            elif val.startswith("✗"):
                f.color.rgb = RED
            elif val.startswith("~"):
                f.color.rgb = AMBER
            else:
                f.color.rgb = INK

# ================================================= SLIDE 3 · TECHNICAL APPROACH
s = S[2]
pointers(s, ["Technologies to be used · Methodology and process for "
             "implementation (Flow Charts/Images/working prototype)"])
steps_a = [("PCAP INGEST", "offline files or\nlive capture"),
           ("TCP REASSEMBLY", "full stream\nreconstruction"),
           ("PROTOCOL ID", "SMTP · IMAP · POP3\nfingerprinting"),
           ("STARTTLS DETECT", "encryption\ntransitions")]
steps_b = [("TLS HANDSHAKE", "version · cipher ·\nkey exchange · FS"),
           ("X.509 VALIDATION", "chain · expiry ·\nkey & signature alg"),
           ("RULES + AI RISK", "deterministic score,\nAI anomaly + advice"),
           ("REPORT + DASHBOARD", "JSON · PDF · HTML,\nself-verifying")]
def flow_row(y, steps):
    x = Inches(0.67); w = Inches(2.78); gap = Inches(0.34)
    for i, (t, d) in enumerate(steps):
        box(s, x, y, w, Inches(1.05), fill=CARD if i % 2 == 0 else INK, line=LINE)
        tc = ACCENT_D if i % 2 == 0 else ACCENT_L
        bc = INK if i % 2 == 0 else RGBColor(0xE4, 0xEE, 0xE8)
        text(s, x + Inches(0.14), y + Inches(0.12), w - Inches(0.28), Inches(0.3),
             t, size=13, bold=True, color=tc)
        text(s, x + Inches(0.14), y + Inches(0.44), w - Inches(0.28), Inches(0.55),
             d.split("\n"), size=10.5, color=bc, space_after=0)
        if i < 3:
            text(s, x + w - Inches(0.02), y + Inches(0.3), gap, Inches(0.4), "→",
                 size=18, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
        x += w + gap
flow_row(Inches(1.6), steps_a)
text(s, Inches(12.35), Inches(2.68), Inches(0.5), Inches(0.4), "↓",
     size=18, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
flow_row(Inches(3.0), steps_b)

# tech chips (left) + architecture image (right)
text(s, Inches(0.67), Inches(4.35), Inches(7), Inches(0.3),
     "TECH STACK — one binary, no cloud dependency", size=13, bold=True,
     color=ACCENT_D)
chips = [("Rust 2024 (unsafe forbidden)", 2.30), ("Tauri 2 desktop", 1.55),
         ("React 18 + TS dashboard", 1.95), ("SQLite + FTS5", 1.30),
         ("Hickory DNS, fail-closed", 1.85), ("sha2 hash-chained audit", 2.10),
         ("WSL2 sandbox", 1.35)]
x = Inches(0.67); y = Inches(4.72)
for label, wch in chips:
    w = Inches(wch)
    if x + w > Inches(7.3):
        x = Inches(0.67); y += Inches(0.52)
    chip(s, x, y, w, Inches(0.4), label, size=10.5)
    x += w + Inches(0.12)
box(s, Inches(0.67), Inches(5.9), Inches(6.55), Inches(0.9),
    fill=RGBColor(0xE9, 0xF6, 0xEC), line=ACCENT)
text(s, Inches(0.85), Inches(6.0), Inches(6.2), Inches(0.75), [[
    ("Endpoint vantage. ", {"bold": True, "color": ACCENT_D, "size": 11.5}),
    ("TLS 1.3 hides certificates from passive sniffers — KIWI is also the "
     "client, so endpoint TlsObservation supplies the ground truth PCAP "
     "cannot see.", {"size": 11.5}),
]])
im = Image.open(ARCH)
ah = Inches(2.3); aw = Emu(int(ah * im.width / im.height))
pic(s, ARCH, Inches(7.7), Inches(4.3), h=ah)
text(s, Inches(7.7), Inches(4.3) + ah + Inches(0.03), aw, Inches(0.28),
     "As-built architecture — Rust workspace of 10 crates",
     size=10, color=MUTED, align=PP_ALIGN.CENTER)

# ================================================= SLIDE 4 · FEASIBILITY
s = S[3]
pointers(s, ["Analysis of the feasibility of the idea · Potential challenges "
             "and risks · Strategies for overcoming these challenges"])
cards = [
    ("FEASIBLE — ALREADY BUILT", ACCENT_D, [
        "PCAP ingest + TCP stream reassembly in kiwi-forensics",
        "TlsObservation capture: version, cipher, FS, cert chain, STARTTLS",
        "Deterministic rule engine → scored, reproducible findings",
        "Self-verifying export envelope (canonical bytes + SHA-256)",
        "Working Tauri client + dashboard; cargo/vitest/CDP CI gates",
    ]),
    ("CHALLENGES & RISKS", AMBER, [
        "TLS 1.3 encrypts certificates — passive capture goes blind",
        "Fragmented / out-of-order PCAP in real networks",
        "AI must never invent findings or verdicts",
        "Heterogeneous server stacks; STARTTLS stripping tricks",
    ]),
    ("OUR MITIGATIONS", ACCENT, [
        "Endpoint vantage fills the passive blind spot — hybrid model",
        "Robust sequencer; incomplete streams flagged, never guessed",
        "Rules are authoritative; AI only classifies + explains",
        "Lab-validated vs. Mailpit/GreenMail corpora + PS dataset",
    ]),
]
x = Inches(0.67)
for title, color, items in cards:
    box(s, x, Inches(1.75), Inches(4.0), Inches(4.7), fill=CARD, line=LINE)
    box(s, x, Inches(1.75), Inches(4.0), Inches(0.62), fill=color)
    text(s, x + Inches(0.2), Inches(1.75), Inches(3.6), Inches(0.62), title,
         size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, x + Inches(0.22), Inches(2.6), Inches(3.6), Inches(3.7),
            [("", i) for i in items], size=12.5, gap=9)
    x += Inches(4.15)
text(s, Inches(0.67), Inches(6.55), Inches(12.2), Inches(0.35), [[
    ("Viable: ", {"bold": True, "color": ACCENT_D, "size": 12.5}),
    ("offline-first desktop — no licences, no cloud, deployable inside "
     "air-gapped SOCs; MPL-2.0 open source.", {"size": 12.5}),
]])

# ================================================= SLIDE 5 · IMPACT
s = S[4]
pointers(s, ["Potential impact on the target audience · Benefits of the "
             "solution (social, economic, etc.)"])
text(s, Inches(0.67), Inches(1.6), Inches(7), Inches(0.35),
     "WHO BENEFITS — AND HOW", size=14, bold=True, color=ACCENT_D)
impacts = [
    ("SOC / DFIR / IR teams", "posture triage in seconds — prioritised findings instead of raw packets"),
    ("Govt. & enterprise admins", "concrete remediation: TLS upgrades, STARTTLS fixes, cert renewals"),
    ("Auditors & compliance", "tamper-evident, self-verifying reports — proof, not claims"),
    ("India / sovereign stack", "offline-capable MPL-2.0 tool — no vendor lock-in, no data egress"),
]
y = Inches(2.05)
for t, d in impacts:
    box(s, Inches(0.67), y, Inches(6.6), Inches(0.95), fill=CARD, line=LINE)
    text(s, Inches(0.9), y + Inches(0.1), Inches(6.2), Inches(0.3), t,
         size=13.5, bold=True, color=ACCENT_D)
    text(s, Inches(0.9), y + Inches(0.42), Inches(6.2), Inches(0.5), d,
         size=11.5, color=INK)
    y += Inches(1.1)
im = Image.open(SHOT_D)
iw = Inches(5.15); ih = Emu(int(iw * im.height / im.width))
pic(s, SHOT_D, Inches(7.55), Inches(1.75), w=iw)
y2 = Inches(1.75) + ih + Inches(0.15)
stats = [("3", "protocols\nSMTP·IMAP·POP3"), ("4", "verdict tiers\nallow→deny"),
         ("1", "self-verifying\nforensic report"), ("0", "fabricated\nfindings")]
x = Inches(7.55)
for n, lab in stats:
    box(s, x, y2, Inches(1.18), Inches(0.95), fill=INK)
    text(s, x, y2 + Inches(0.08), Inches(1.18), Inches(0.4), n, size=20,
         bold=True, color=ACCENT_L, align=PP_ALIGN.CENTER)
    text(s, x + Inches(0.05), y2 + Inches(0.5), Inches(1.08), Inches(0.45),
         lab.split("\n"), size=8.5, color=RGBColor(0xC9, 0xD6, 0xCE),
         align=PP_ALIGN.CENTER, space_after=0)
    x += Inches(1.34)

# ================================================= SLIDE 6 · REFERENCES
s = S[5]
pointers(s, ["Details / Links of the reference and research work"])
text(s, Inches(0.67), Inches(1.7), Inches(6), Inches(0.35),
     "STANDARDS & GUIDANCE", size=14, bold=True, color=ACCENT_D)
bullets(s, Inches(0.67), Inches(2.1), Inches(6.0), Inches(4.4), [
    ("RFC 3207 / 2595", "SMTP & IMAP STARTTLS negotiation"),
    ("RFC 8314", "cleartext mail deprecated — implicit TLS"),
    ("RFC 8461 / 8996", "MTA-STS · TLS 1.0/1.1 deprecated"),
    ("NIST SP 800-52r2", "TLS configuration guidelines"),
    ("CA/B Forum BRs", "certificate validation baseline"),
    ("MITRE ATT&CK", "T1040 sniffing · T1557 adversary-in-the-middle"),
], size=12.5, gap=8)
text(s, Inches(7.1), Inches(1.7), Inches(5.8), Inches(0.35),
     "BASELINES & PROJECT EVIDENCE", size=14, bold=True, color=ACCENT_D)
bullets(s, Inches(7.1), Inches(2.1), Inches(5.6), Inches(4.4), [
    ("Wireshark · Zeek", "packet tooling baseline — decode, no posture"),
    ("SSLyze · testssl.sh", "active-scan baseline compared in slide 2"),
    ("SIH26159 dataset", "attached synthetic PCAP corpus for validation"),
    ("KIWI repo", "docs/contracts/forensics.md — FSV-1 report vocabulary"),
    ("Working prototype", "10-crate Rust workspace, CI-gated, ~30 MB binary"),
], size=12.5, gap=8)
box(s, Inches(0.67), Inches(5.3), Inches(12.28), Inches(1.3), fill=INK)
text(s, Inches(1.0), Inches(5.52), Inches(11.6), Inches(0.9), [
    [("KIWI doesn't ask you to trust it — it shows you the evidence.",
      {"size": 18, "bold": True, "color": ACCENT_L})],
    [("Every claim in this deck is verifiable in the repo: "
      "deterministic capture, rule-based scoring, hash-chained audit, "
      "self-verifying exports.", {"size": 12.5,
                                  "color": RGBColor(0xC9, 0xD6, 0xCE)})],
], space_after=4)

# ---- delete slide 7 (instructions) per template note --------------------
xml_slides = prs.slides._sldIdLst
slides_list = list(xml_slides)
xml_slides.remove(slides_list[6])

prs.save(OUT)
print("wrote", OUT)
